from playwright.sync_api import sync_playwright
import os
import re
from openpyxl import Workbook
from datetime import datetime

def extrair_numeros_processos_pagina(page):
    """Extrai números de processo da página atual (lista de processos)."""
    links = page.locator("a").all()
    numeros = []
    padrao = re.compile(r'\d{5}\.\d{6}/\d{4}-\d{2}')
    for link in links:
        texto = link.inner_text().strip()
        if padrao.search(texto):
            numeros.append(texto)
    return numeros

def obter_total_paginas(page):
    """Obtém o número total de páginas a partir do texto '(1 de 9)'."""
    try:
        paginator = page.locator("#tblProcessosMovimentacao_paginator_top .ui-paginator-current")
        texto = paginator.inner_text()
        match = re.search(r'de\s+(\d+)', texto)
        if match:
            total = int(match.group(1))
            print(f"📄 Total de páginas detectado: {total}")
            return total
    except Exception as e:
        print(f"⚠️ Não foi possível detectar total de páginas: {e}")
    return 1

def navegar_para_pagina(page, numero_pagina):
    """Navega diretamente para uma página clicando no botão numérico."""
    try:
        btn_pagina = page.locator(f"#tblProcessosMovimentacao_paginator_top .ui-paginator-page:has-text('{numero_pagina}')")
        if btn_pagina.count() > 0:
            btn_pagina.click()
            page.wait_for_load_state("networkidle")
            page.wait_for_timeout(3000)
            return True
    except Exception as e:
        print(f"⚠️ Erro ao navegar para página {numero_pagina}: {e}")
    return False

def coletar_todos_processos(page):
    """Navega por todas as páginas e coleta todos os números de processo."""
    todos_numeros = []
    total_paginas = obter_total_paginas(page)
    
    print(f"🔍 Coletando processos de {total_paginas} página(s)...")
    
    for pagina in range(1, total_paginas + 1):
        if pagina > 1:
            if not navegar_para_pagina(page, pagina):
                print(f"⚠️ Não foi possível navegar para página {pagina}. Parando coleta.")
                break
        numeros_pag = extrair_numeros_processos_pagina(page)
        todos_numeros.extend(numeros_pag)
        print(f"Página {pagina}: encontrados {len(numeros_pag)} processos. Total parcial: {len(todos_numeros)}")
        page.wait_for_timeout(1000)
    
    print(f"\n📊 Total coletado: {len(todos_numeros)} processos em {total_paginas} página(s).")
    return todos_numeros

def fazer_busca_processo(page, numero_processo):
    """Preenche o campo de busca rápida com o número do processo e submete."""
    try:
        campo_busca = page.locator("#formBuscaRapida\\:buscaRapida")
        campo_busca.wait_for(state="visible", timeout=0)
        campo_busca.fill(numero_processo)
        botao_busca = page.locator("#formBuscaRapida\\:btBuscaRapida")
        botao_busca.click()
        page.wait_for_selector("[id='formProcesso:incluirDocumentoRequerimento']", state="visible", timeout=0)
        print(f"✅ Carregou processo: {numero_processo}")
        return True
    except Exception as e:
        print(f"❌ Falha ao buscar processo {numero_processo}: {e}")
        return False

def processar_fluxo_processo(page, numero_processo, download_path):
    """Executa todo o fluxo para um único processo."""
    try:
        incluir_doc = page.locator("[id='formProcesso:incluirDocumentoRequerimento']")
        incluir_doc.wait_for(state="visible", timeout=0)
        incluir_doc.click()
        print(f"✅ [{numero_processo}] Clicou em 'Incluir Documento do Requerimento'")
        
        page.wait_for_selector("[id='formProcesso:dlgAnexarDocumentoRequerimento']", state="visible", timeout=0)
        page.wait_for_timeout(1000)
        
        try:
            radio_label = page.get_by_text("Incluir o Documento da Tarefa de Recurso Ordinário", exact=True)
            radio_label.wait_for(state="visible", timeout=0)
            radio_label.click()
            print(f"✅ [{numero_processo}] Opção selecionada via texto")
        except:
            radio_input = page.locator("[id='formProcesso:sltDocumentoRequerimento:0']")
            radio_input.wait_for(state="visible", timeout=0)
            radio_input.check()
            print(f"✅ [{numero_processo}] Opção via radio input")
        
        page.wait_for_timeout(3000)
        
        try:
            confirmar_btn = page.locator("[id='formProcesso:btConfirmarAnexarDocumentoRequerimento']")
            confirmar_btn.wait_for(state="visible", timeout=0)
            confirmar_btn.click()
            print(f"✅ [{numero_processo}] Confirmou")
        except:
            page.evaluate("""() => { document.querySelector("[id='formProcesso:btConfirmarAnexarDocumentoRequerimento']").click(); }""")
            print(f"✅ [{numero_processo}] Confirmou via JS")
        
        page.wait_for_timeout(3000)
        
        arvore_doc = page.locator("[id='formProcesso:arvoreDocumental']")
        arvore_doc.wait_for(state="visible", timeout=0)
        arvore_doc.click()
        print(f"✅ [{numero_processo}] Clicou em 'Árvore Documental'")
        page.wait_for_timeout(3000)
        
        try:
            checkbox_tudo = page.locator(".ui-chkbox-box").first
            checkbox_tudo.wait_for(state="visible", timeout=0)
            checkbox_tudo.click()
            print(f"✅ [{numero_processo}] Selecionar tudo")
        except:
            page.get_by_text("Selecionar tudo").click()
            print(f"✅ [{numero_processo}] Selecionar tudo via texto")
        
        page.wait_for_timeout(3000)
        
        try:
            btn_pdf = page.locator("[id='arvoreDocumentosForm:btnGerarPdf']")
            btn_pdf.wait_for(state="visible", timeout=0)
            btn_pdf.click()
            print(f"✅ [{numero_processo}] Gerar PDF")
        except:
            page.get_by_text("Gerar PDF").click()
            print(f"✅ [{numero_processo}] Gerar PDF via texto")
        
        print(f"⏳ [{numero_processo}] Aguardando visualizador PDF...")
        page.wait_for_selector("[id='arvoreDocumentosForm:dlgVisualizadorDoc']", state="visible", timeout=0)
        page.wait_for_timeout(2000)
        
        print(f"⏳ [{numero_processo}] Aguardando carregamento completo do PDF...")
        iframe_locator = page.frame_locator("[id='arvoreDocumentosForm:j_idt118']")
        
        try:
            num_pages_locator = iframe_locator.locator("#numPages")
            num_pages_locator.wait_for(state="visible", timeout=0)
            last_text = ""
            stable_count = 0
            for _ in range(20):
                text = num_pages_locator.inner_text()
                if text != last_text:
                    last_text = text
                    stable_count = 0
                else:
                    stable_count += 1
                if stable_count >= 2 and "de" in text and len(text) > 3:
                    match = re.search(r'\d+', text)
                    if match and int(match.group()) > 0:
                        print(f"✅ [{numero_processo}] PDF carregado: {text} páginas")
                        break
                page.wait_for_timeout(500)
            page.wait_for_timeout(2000)
        except Exception as e:
            print(f"⚠️ [{numero_processo}] Não foi possível detectar número de páginas: {e}")
        
        print(f"🔍 [{numero_processo}] Reduzindo zoom para 80%...")
        for _ in range(2):
            page.keyboard.press("Control+-")
        page.wait_for_timeout(2000)
        
        try:
            download_btn = iframe_locator.locator("#download").first
            download_btn.wait_for(state="visible", timeout=0)
            iframe_locator.locator("body").evaluate("el => el.scrollTop = el.scrollHeight")
            page.wait_for_timeout(500)
            with page.expect_download() as download_info:
                download_btn.click()
            download = download_info.value
            suggested_filename = download.suggested_filename or "documento.pdf"
            safe_num = numero_processo.replace("/", "_").replace("\\", "_")
            target_path = os.path.join(download_path, f"{safe_num}_{suggested_filename}")
            download.save_as(target_path)
            print(f"✅ [{numero_processo}] PDF baixado: {target_path}")
            page.keyboard.press("Control+0")
            return True
        except Exception as e:
            print(f"❌ [{numero_processo}] Falha no download: {e}")
            page.keyboard.press("Control+0")
            return False
    except Exception as e:
        print(f"❌ [{numero_processo}] Erro durante fluxo: {e}")
        return False

def main():
    download_path = r"C:\Users\Pedro\Downloads\processos"
    os.makedirs(download_path, exist_ok=True)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Processos"
    ws.append(["Processo", "Página", "Status", "Data/Hora"])
    
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        context = browser.new_context(accept_downloads=True, viewport={"width": 1920, "height": 1080})
        page = context.new_page()
        
        page.goto("https://esisrec.inss.gov.br/esisrec/")
        
        print("=" * 60)
        print("Faça o login manualmente no navegador.")
        print("O robô aguardará até a tela principal aparecer.")
        print("=" * 60)
        
        page.wait_for_url(lambda url: "esisrec.inss.gov.br/esisrec/" in url, timeout=0)
        print("✅ Redirecionamento pós-login detectado.")
        
        try:
            unidade_select = page.wait_for_selector("select#domains", timeout=10000)
            if unidade_select:
                print("🔍 Tela de unidade detectada. Selecionando 'GOV:CRPS.10ª JR'...")
                page.select_option("select#domains", label="GOV:CRPS.10ª JR")
                page.wait_for_load_state("networkidle")
                page.wait_for_timeout(2000)
                print("✅ Unidade selecionada.")
        except:
            print("ℹ️ Nenhuma tela de unidade.")
        
        print("⏳ Aguardando o link 'Processos com Localizador Relatar'...")
        link_relatar = page.get_by_text("Processos com Localizador \"Relatar\"", exact=True)
        link_relatar.wait_for(state="visible", timeout=0)
        link_relatar.click()
        print("✅ Clicou em 'Relatar' - carregando lista de processos...")
        page.wait_for_timeout(5000)
        
        # Coleta todos os processos de todas as páginas
        todos_numeros = coletar_todos_processos(page)
        
        if not todos_numeros:
            print("⚠️ Nenhum processo encontrado. Encerrando.")
            return
        
        for num in todos_numeros:
            ws.append([num, "N/A", "Pendente", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        wb.save(os.path.join(download_path, "processos_log.xlsx"))
        
        for idx, numero in enumerate(todos_numeros):
            print(f"\n--- Processando {idx+1}/{len(todos_numeros)}: {numero} ---")
            
            try:
                btn_inicio = page.locator("[id='formProcesso:btPaginaInicialFim']")
                if btn_inicio.count() == 0:
                    btn_inicio = page.locator("[id='formProcesso:btPaginaInicial']")
                btn_inicio.wait_for(state="visible", timeout=10000)
                btn_inicio.click()
                print("✅ Clicou em 'Página Inicial'")
                page.wait_for_timeout(3000)
            except:
                page.goto("https://esisrec.inss.gov.br/esisrec/")
                page.wait_for_timeout(3000)
            
            if not fazer_busca_processo(page, numero):
                ws.cell(row=idx+2, column=3, value="Erro na busca")
                ws.cell(row=idx+2, column=4, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
                wb.save(os.path.join(download_path, "processos_log.xlsx"))
                continue
            
            sucesso = processar_fluxo_processo(page, numero, download_path)
            status = "Baixado" if sucesso else "Erro no fluxo"
            ws.cell(row=idx+2, column=3, value=status)
            ws.cell(row=idx+2, column=4, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            wb.save(os.path.join(download_path, "processos_log.xlsx"))
            page.wait_for_timeout(2000)
        
        print("\n🎉 Todos os processos foram processados!")
        print(f"Planilha salva em: {download_path}\\processos_log.xlsx")
        print("O navegador permanecerá aberto. Feche-o quando terminar.")
        
        try:
            page.wait_for_event("close")
        except:
            pass
        finally:
            context.close()
            browser.close()

if __name__ == "__main__":
    main()